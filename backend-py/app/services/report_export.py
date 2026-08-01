import io
import json
from datetime import datetime
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app.services.application_status import explain_application_status
from app.services.lead_scoring import (
    ACADEMIC_BACKGROUND_LABELS,
    MIGRATION_INTENT_LABELS,
    PREFERRED_COUNTRY_ANY,
)

_BRAND_COLOR = "1a5f7a"


def _build_report_data(row: Any) -> dict:
    """Deliberately excludes internal CRM-only fields (Hot/Warm/Cold score,
    counsellor notes, assignment, follow-up data) — this document is the
    student's own copy of their recommendation, not a CRM record."""
    status_info = explain_application_status(row["application_status"])
    academic_background = row["academic_background"]
    migration_intent = row["migration_intent"]
    preferred_country = row["preferred_country"]

    return {
        "name": row["name"],
        "email": row["email"],
        "phone": row["phone"],
        "referenceCode": row["reference_code"],
        "generatedAt": datetime.now().strftime("%B %d, %Y"),
        "profile": [
            {"label": "GPA", "value": f"{row['gpa']} / 4.0"},
            {"label": "IELTS", "value": f"{row['ielts']} / 9.0"},
            {"label": "Annual Budget", "value": f"${row['budget']:,.0f} USD"},
            {"label": "Study Gap", "value": f"{row['gap']} year(s)"},
            {
                "label": "Academic Background",
                "value": ACADEMIC_BACKGROUND_LABELS.get(academic_background, academic_background) if academic_background else "—",
            },
            {"label": "Career Goals", "value": row["career_goals"] or "—"},
            {
                "label": "Migration Intent",
                "value": MIGRATION_INTENT_LABELS.get(migration_intent, migration_intent) if migration_intent else "—",
            },
            {
                "label": "Preferred Country",
                "value": preferred_country if preferred_country and preferred_country != PREFERRED_COUNTRY_ANY else "Any / Let AI decide",
            },
        ],
        "recommendedCountries": row["countries"] or "—",
        "aiResponse": row["ai_response"] or "(No recommendation on file yet.)",
        "nextSteps": json.loads(row["next_steps"]) if row["next_steps"] else [],
        "applicationStatus": {"label": status_info.label, "description": status_info.description},
    }


def _wrap_text(text: str, width_chars: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        words = paragraph.split(" ")
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if len(candidate) > width_chars and current:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    return lines


def generate_report_pdf(row: Any) -> bytes:
    data = _build_report_data(row)
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter
    margin = 50
    y = height - margin

    def heading(text: str) -> None:
        nonlocal y
        c.setFont("Helvetica-Bold", 12)
        c.setFillColorRGB(0x1A / 255, 0x5F / 255, 0x7A / 255)
        c.drawString(margin, y, text)
        c.setFillColorRGB(0, 0, 0)
        y -= 18

    def line(text: str, font: str = "Helvetica", size: int = 10, gap: int = 14) -> None:
        nonlocal y
        c.setFont(font, size)
        for wrapped in _wrap_text(text, 95):
            if y < margin + 40:
                c.showPage()
                y = height - margin
                c.setFont(font, size)
            c.drawString(margin, y, wrapped)
            y -= gap

    c.setFont("Helvetica-Bold", 20)
    c.setFillColorRGB(0x1A / 255, 0x5F / 255, 0x7A / 255)
    c.drawString(margin, y, "AIEC Global")
    y -= 20
    c.setFont("Helvetica", 13)
    c.setFillColorRGB(0.33, 0.33, 0.33)
    c.drawString(margin, y, "Study Abroad Recommendation Report")
    y -= 16
    c.setFont("Helvetica", 9)
    c.setFillColorRGB(0.6, 0.6, 0.6)
    ref_suffix = f"  ·  Reference: {data['referenceCode']}" if data["referenceCode"] else ""
    c.drawString(margin, y, f"Generated {data['generatedAt']}{ref_suffix}")
    y -= 24
    c.setFillColorRGB(0, 0, 0)

    c.setFont("Helvetica-Bold", 14)
    c.drawString(margin, y, data["name"])
    y -= 16
    contact_line = "  ·  ".join([v for v in [data["email"], data["phone"]] if v])
    if contact_line:
        c.setFont("Helvetica", 10)
        c.setFillColorRGB(0.33, 0.33, 0.33)
        c.drawString(margin, y, contact_line)
        y -= 14
    c.setFillColorRGB(0, 0, 0)
    y -= 10

    heading("Your Profile")
    for item in data["profile"]:
        line(f"{item['label']}: {item['value']}")
    y -= 10

    heading("Recommended Countries")
    line(data["recommendedCountries"])
    y -= 10

    heading("AI Recommendation")
    line(data["aiResponse"])
    y -= 10

    if data["nextSteps"]:
        heading("Next Steps")
        for i, step in enumerate(data["nextSteps"], start=1):
            line(f"{i}. {step}")
        y -= 10

    heading("Current Application Status")
    line(f"{data['applicationStatus']['label']} — {data['applicationStatus']['description']}")

    y -= 20
    c.setFont("Helvetica", 8)
    c.setFillColorRGB(0.6, 0.6, 0.6)
    c.drawString(margin, y, "Questions about this report? Contact your AIEC Global counsellor.")

    c.save()
    return buffer.getvalue()


def generate_report_docx(row: Any) -> bytes:
    data = _build_report_data(row)
    doc = Document()
    brand = RGBColor(0x1A, 0x5F, 0x7A)
    grey = RGBColor(0x55, 0x55, 0x55)
    light_grey = RGBColor(0x99, 0x99, 0x99)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = brand

    title = doc.add_paragraph()
    run = title.add_run("AIEC Global")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = brand

    subtitle = doc.add_paragraph()
    run = subtitle.add_run("Study Abroad Recommendation Report")
    run.font.size = Pt(13)
    run.font.color.rgb = grey

    meta = doc.add_paragraph()
    ref_suffix = f"  ·  Reference: {data['referenceCode']}" if data["referenceCode"] else ""
    run = meta.add_run(f"Generated {data['generatedAt']}{ref_suffix}")
    run.font.size = Pt(9)
    run.font.color.rgb = light_grey

    name_p = doc.add_paragraph()
    run = name_p.add_run(data["name"])
    run.bold = True
    run.font.size = Pt(14)

    contact_line = "  ·  ".join([v for v in [data["email"], data["phone"]] if v])
    if contact_line:
        contact_p = doc.add_paragraph()
        run = contact_p.add_run(contact_line)
        run.font.size = Pt(10)
        run.font.color.rgb = grey

    heading("Your Profile")
    for item in data["profile"]:
        p = doc.add_paragraph()
        r1 = p.add_run(f"{item['label']}: ")
        r1.bold = True
        p.add_run(item["value"])

    heading("Recommended Countries")
    doc.add_paragraph(data["recommendedCountries"])

    heading("AI Recommendation")
    for para in data["aiResponse"].split("\n"):
        doc.add_paragraph(para)

    if data["nextSteps"]:
        heading("Next Steps")
        for i, step in enumerate(data["nextSteps"], start=1):
            doc.add_paragraph(f"{i}. {step}")

    heading("Current Application Status")
    p = doc.add_paragraph()
    r1 = p.add_run(data["applicationStatus"]["label"])
    r1.bold = True
    p.add_run(f" — {data['applicationStatus']['description']}")

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(20)
    run = footer.add_run("Questions about this report? Contact your AIEC Global counsellor.")
    run.font.size = Pt(8)
    run.font.color.rgb = light_grey

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
